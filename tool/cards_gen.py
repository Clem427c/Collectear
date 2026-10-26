import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

rng = random.Random(2026)

RARITIES = ["Commune", "Peu commune", "Rare", "Épique", "Légendaire"]
BUDGET = {"Commune": 30, "Peu commune": 37, "Rare": 45, "Épique": 55, "Légendaire": 68}
DROP = {"Commune": 0.60, "Peu commune": 0.25, "Rare": 0.10, "Épique": 0.04, "Légendaire": 0.01}
COUNT_CREA = {"Commune": 80, "Peu commune": 60, "Rare": 34, "Épique": 18, "Légendaire": 8}
COUNT_HERO = {"Commune": 40, "Peu commune": 30, "Rare": 16, "Épique": 10, "Légendaire": 4}
PROFILES = {  # parts du budget : PV, ATQ, DEF
    "Équilibré": (0.34, 0.33, 0.33),
    "Offensif": (0.25, 0.50, 0.25),
    "Défensif": (0.30, 0.20, 0.50),
    "Robuste": (0.55, 0.25, 0.20),
}
PV_COEF = 5
JITTER = 0.03

# ---------- Noms ----------
LEG_CREA = ["Astrelune", "Volcarion", "Abyssiel", "Tempestor", "Sylvarche", "Glacivore", "Chronoptère", "Néantor"]
LEG_HERO = ["Aurèle l'Immortelle", "Kaelith, Gardien des Seuils", "Sélénor le Dernier Mage", "Ysolde Brise-Étoiles"]
EPI_CREA = ["Pyralis", "Hydromance", "Fulgurion", "Basaltor", "Noctifère", "Cristalame", "Orageon", "Mirelune", "Sabrecorne",
            "Ronceroi", "Cendrolithe", "Aquilarc", "Vortexis", "Sépulcrin", "Émeraudon", "Marévore", "Solsticien", "Givrespectre"]
EPI_HERO = ["Thaïs, Lame d'Aube", "Bram Cœur-de-Fer", "Liora la Tisseuse d'Orage", "Oren le Chasseur d'Ombres",
            "Maëlis l'Archère d'Argent", "Doran le Briseur de Rocs", "Isaure la Prêtresse des Marées",
            "Vesper le Voleur de Lunes", "Calixte le Chevalier Errant", "Nahia la Dompteuse de Bêtes"]

STEMS = ["Brais", "Givr", "Mouss", "Brum", "Écum", "Sabl", "Cendr", "Grel", "Silex", "Saul", "Nacr", "Ambr",
         "Ombr", "Tonn", "Flamm", "Glac", "Roch", "Pétal", "Plum", "Croc", "Griff", "Écaill", "Lun", "Vent",
         "Coquill", "Épin", "Torr", "Lich", "Gland", "Fang", "Bulb", "Cristall", "Rouill", "Pomm", "Tourb", "Ronc",
         "Lierr", "Chard", "Myrt", "Charb", "Souffr", "Quartz", "Opal", "Jasp", "Saph", "Galet", "Dun", "Marr",
         "Nois", "Trembl", "Guêp", "Frel", "Léz", "Grenouill", "Têt", "Hib", "Corb", "Loutr", "Blair", "Hériss",
         "Taup", "Mul", "Belett", "Cerv", "Lynx", "Ours", "Sangl", "Salam", "Trit", "Crap", "Escarg", "Libell",
         "Papill", "Abeill", "Més", "Pins", "Cigal", "Geck", "Fourm", "Crabb", "Mouett", "Pieuvr", "Rasc", "Gob"]
SUF = {
    "Commune": ["ou", "iot", "ette", "in", "ichon", "inou"],
    "Peu commune": ["ard", "aud", "erin", "ille", "ouille", "eron"],
    "Rare": ["ix", "orn", "axe", "idor", "anthe", "oros"],
}
EXCLUDE = {"Herbette", "Plumette", "Sablon", "Rochette", "Pommette", "Lunette", "Glandin", "Galetou", "Moussard",
           "Mouettette", "Belettette", "Charbin", "Rouillard", "Brumaud", "Crapaud", "Crapette", "Coquillard", "Galetin",
           "Pinsin", "Gobin", "Gobette", "Tritin", "Mulet", "Marron", "Ourson", "Lynxette", "Têtard", "Corbin", "Taupin",
           "Lichette", "Saulille", "Sanglier", "Fourmille"}

def creature_names(rarity, n, used):
    stems = STEMS[:]
    rng.shuffle(stems)
    out = []
    for st in stems:
        opts = [x for x in SUF[rarity]
                if (st + x) not in EXCLUDE and (st + x) not in used
                and not (st.endswith("ill") and ("ill" in x or x.startswith("i")))
                and not (st.endswith("tt") and x == "ette")]
        if not opts:
            continue
        name = st + rng.choice(opts)
        out.append(name); used.add(name)
        if len(out) == n:
            return out
    raise SystemExit(f"pas assez de racines pour {rarity}: {len(out)}")

F_NAMES = ["Aëla", "Brune", "Céleste", "Dalia", "Elna", "Faye", "Gwenn", "Hélia", "Ilona", "Jade", "Kira", "Léonie", "Mira",
           "Nola", "Opale", "Prune", "Rosa", "Sienne", "Talia", "Uma", "Vaïna", "Wenna", "Yara", "Zélie", "Alba", "Bérénice",
           "Cyrielle", "Doria", "Elvire", "Flora", "Gaïa", "Hortense", "Inès", "Jouna", "Lysa", "Mélusine", "Néva", "Ondine",
           "Perle", "Rhéa", "Salomé", "Tessa", "Ulla", "Viviane", "Wanda", "Ysa", "Zoé", "Anouk", "Capucine", "Elsa"]
M_NAMES = ["Aldo", "Basile", "Cyprien", "Dorian", "Elio", "Félix", "Gaspard", "Hugo", "Ilan", "Jory", "Kaël", "Lino", "Milo",
           "Noé", "Octave", "Pio", "Rémi", "Silas", "Tobias", "Ugo", "Victor", "Wilfrid", "Yanis", "Zéphyr", "Achille",
           "Bastien", "Clovis", "Darius", "Ernest", "Firmin", "Gédéon", "Hector", "Isidore", "Joachim", "Léandre", "Marius",
           "Nils", "Oscar", "Pacôme", "Raphaël", "Sacha", "Théo", "Ulysse", "Valère", "Wilhelm", "Ewen", "Anatole",
           "Barnabé", "Corentin", "Edmond"]
ROLES = {
    "Commune": [("le Vagabond", "la Vagabonde"), ("l'Apprenti", "l'Apprentie"), ("le Pêcheur", "la Pêcheuse"),
                ("le Bûcheron", "la Bûcheronne"), ("le Garde", "la Garde"), ("l'Éclaireur", "l'Éclaireuse"),
                ("le Forgeron", "la Forgeronne"), ("le Berger", "la Bergère"), ("le Cueilleur", "la Cueilleuse"),
                ("le Marchand", "la Marchande")],
    "Peu commune": [("le Rôdeur", "la Rôdeuse"), ("le Mercenaire", "la Mercenaire"), ("l'Alchimiste", "l'Alchimiste"),
                    ("le Duelliste", "la Duelliste"), ("le Pisteur", "la Pisteuse"), ("le Moine", "la Moniale"),
                    ("le Corsaire", "la Corsaire"), ("le Barde", "la Barde")],
    "Rare": [("le Paladin", "la Paladine"), ("le Mage des Cendres", "la Mage des Cendres"), ("le Maître-Lame", "la Maîtresse-Lame"),
             ("le Druide", "la Druidesse"), ("l'Arcaniste", "l'Arcaniste"), ("le Gardien des Ruines", "la Gardienne des Ruines"),
             ("le Chasseur de Primes", "la Chasseuse de Primes"), ("le Chevaucheur de Wyverne", "la Chevaucheuse de Wyverne")],
}
first_pool = [(n, "f") for n in F_NAMES] + [(n, "m") for n in M_NAMES]
rng.shuffle(first_pool)

def hero_names(rarity, n):
    out = []
    roles = ROLES[rarity]
    for k in range(n):
        name, g = first_pool.pop()
        role = roles[k % len(roles)][0 if g == "m" else 1]
        out.append(f"{name} {role}")
    return out

# ---------- Cartes ----------
cards, used = [], set(LEG_CREA + EPI_CREA)
for r in RARITIES:
    if r == "Légendaire":
        cn, hn = LEG_CREA, LEG_HERO
    elif r == "Épique":
        cn, hn = EPI_CREA, EPI_HERO
    else:
        cn, hn = creature_names(r, COUNT_CREA[r], used), hero_names(r, COUNT_HERO[r])
    assert len(cn) == COUNT_CREA[r] and len(hn) == COUNT_HERO[r], r
    group = [(n, "Créature") for n in cn] + [(n, "Héros") for n in hn]
    rng.shuffle(group)
    plist = list(PROFILES)
    for idx, (name, fam) in enumerate(group):
        prof = plist[idx % 4]
        base = PROFILES[prof]
        parts = [max(0.12, b + rng.uniform(-0.03, 0.03)) for b in base]
        t = sum(parts); parts = [round(p / t, 3) for p in parts]
        parts[2] = round(1 - parts[0] - parts[1], 3)
        var = round(1 + rng.uniform(-JITTER, JITTER), 3)
        cards.append([name, fam, r, prof, *parts, var])

RENAME = {"Galetette": "Galetiot"}
for c in cards:
    c[0] = RENAME.get(c[0], c[0])
names = [c[0] for c in cards]
assert len(cards) == 300 and len(set(names)) == 300

# ---------- Classeur ----------
wb = Workbook()
FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1B1440")
BLUE = Font(name=FONT, color="0000FF")
BODY = Font(name=FONT)
B = Font(name=FONT, bold=True)
thin = Side(style="thin", color="D9D6E8")
RCOL = {"Commune": "EDEDED", "Peu commune": "DDF5F2", "Rare": "E7E1FF", "Épique": "FFE3E3", "Légendaire": "FFF0CC"}

# Règles
ru = wb.active; ru.title = "Règles"
ru["A1"] = "Collectear – règles d'équilibrage des cartes"; ru["A1"].font = Font(name=FONT, bold=True, size=14)
ru["A2"] = "Les cellules en bleu sont modifiables : toutes les stats de l'onglet Cartes se recalculent."; ru["A2"].font = Font(name=FONT, italic=True, color="6B6880")
hdr = ["Rareté", "Nb de cartes", "Budget de puissance", "Taux d'obtention (GDD)"]
for j, h in enumerate(hdr, 1):
    c = ru.cell(row=4, column=j, value=h); c.font = H; c.fill = HF
for i, r in enumerate(RARITIES):
    row = 5 + i
    ru.cell(row=row, column=1, value=r).font = B
    ru.cell(row=row, column=1).fill = PatternFill("solid", fgColor=RCOL[r])
    ru.cell(row=row, column=2, value=f"=COUNTIF(Cartes!$D:$D,A{row})").font = BODY
    ru.cell(row=row, column=3, value=BUDGET[r]).font = BLUE
    c = ru.cell(row=row, column=4, value=DROP[r]); c.font = BLUE; c.number_format = "0%"
ru["A10"] = "Total"; ru["A10"].font = B
ru["B10"] = "=SUM(B5:B9)"; ru["B10"].font = B
ru["A12"] = "Coefficient PV"; ru["A12"].font = B
ru["C12"] = PV_COEF; ru["C12"].font = BLUE
ru["A13"] = "Variation max par carte"; ru["A13"].font = B
ru["C13"] = JITTER; ru["C13"].font = BLUE; ru["C13"].number_format = "0%"
ru["C13"].comment = Comment("Valeur utilisée lors de la génération des colonnes « Variation » ; la modifier ici ne régénère pas ces colonnes.", "Claude")

ru["A15"] = "Formules"; ru["A15"].font = Font(name=FONT, bold=True, size=12)
lines = [
    "Puissance = PV / Coefficient PV + ATQ + DEF  (1 point de puissance = 5 PV = 1 ATQ = 1 DEF)",
    "PV = arrondi(Budget × Variation × Part PV × Coefficient PV)",
    "ATQ = arrondi(Budget × Variation × Part ATQ)        DEF = arrondi(Budget × Variation × Part DEF)",
    "Chaque palier de rareté a environ +22 % de budget sur le précédent ; avec ±3 % de variation, les plages ne se chevauchent jamais.",
    "Le profil répartit le budget entre les stats ; chaque carte a une petite variation aléatoire de ses parts (±3 points).",
]
for k, t in enumerate(lines):
    ru.cell(row=16 + k, column=1, value=t).font = BODY

ru["A22"] = "Profils (parts de base du budget)"; ru["A22"].font = Font(name=FONT, bold=True, size=12)
for j, h in enumerate(["Profil", "Part PV", "Part ATQ", "Part DEF"], 1):
    c = ru.cell(row=23, column=j, value=h); c.font = H; c.fill = HF
for i, (p, v) in enumerate(PROFILES.items()):
    ru.cell(row=24 + i, column=1, value=p).font = B
    for j, x in enumerate(v):
        c = ru.cell(row=24 + i, column=2 + j, value=x); c.font = BODY; c.number_format = "0%"
ru["A28"] = "Source : budgets, profils et répartition 120 / 90 / 50 / 28 / 12 proposés par Claude ; taux d'obtention repris du GDD (exemple, à équilibrer en playtest)."
ru["A28"].font = Font(name=FONT, italic=True, color="6B6880", size=9)
for col, w in zip("ABCD", [26, 14, 20, 22]):
    ru.column_dimensions[col].width = w

# Cartes
ws = wb.create_sheet("Cartes")
heads = ["N°", "Nom", "Famille", "Rareté", "Profil", "PV", "ATQ", "DEF", "Puissance", "Part PV", "Part ATQ", "Part DEF", "Variation"]
for j, h in enumerate(heads, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center")
bud = "INDEX(Règles!$C$5:$C$9,MATCH($D{r},Règles!$A$5:$A$9,0))"
for i, (name, fam, r, prof, pp, pa, pd, var) in enumerate(cards):
    row = i + 2
    b = bud.format(r=row)
    vals = [i + 1, name, fam, r, prof,
            f"=ROUND({b}*$M{row}*$J{row}*Règles!$C$12,0)",
            f"=ROUND({b}*$M{row}*$K{row},0)",
            f"=ROUND({b}*$M{row}*$L{row},0)",
            f"=F{row}/Règles!$C$12+G{row}+H{row}",
            pp, pa, pd, var]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v); c.font = BODY
        c.border = Border(bottom=thin)
    ws.cell(row=row, column=1).number_format = "000"
    ws.cell(row=row, column=4).fill = PatternFill("solid", fgColor=RCOL[r])
    ws.cell(row=row, column=9).number_format = "0.0"
    for j in (10, 11, 12):
        ws.cell(row=row, column=j).number_format = "0.0%"; ws.cell(row=row, column=j).font = Font(name=FONT, color="6B6880")
    ws.cell(row=row, column=13).number_format = "0.000"; ws.cell(row=row, column=13).font = Font(name=FONT, color="6B6880")
last = len(cards) + 1
tab = Table(displayName="Cartes", ref=f"A1:M{last}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(tab)
for c, w in zip("ABCDEFGHIJKLM", [6, 36, 11, 14, 12, 7, 7, 7, 11, 9, 9, 9, 10]):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "C2"
ws["J1"].comment = Comment("Colonnes J à M : paramètres de génération de chaque carte (parts du budget et variation). Les stats en découlent.", "Claude")

# Synthèse
sy = wb.create_sheet("Synthèse")
sy["A1"] = "Synthèse par rareté"; sy["A1"].font = Font(name=FONT, bold=True, size=14)
sh = ["Rareté", "Nb", "Créatures", "Héros", "PV moyen", "ATQ moyenne", "DEF moyenne", "Puissance min", "Puissance moy.", "Puissance max", "Marge sur rareté inférieure"]
for j, h in enumerate(sh, 1):
    c = sy.cell(row=3, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center", wrap_text=True)
R = "Cartes!$D$2:$D$%d" % last
for i, r in enumerate(RARITIES):
    row = 4 + i
    sy.cell(row=row, column=1, value=r).font = B
    sy.cell(row=row, column=1).fill = PatternFill("solid", fgColor=RCOL[r])
    f = {
        2: f"=COUNTIF({R},A{row})",
        3: f'=COUNTIFS({R},A{row},Cartes!$C$2:$C${last},"Créature")',
        4: f'=COUNTIFS({R},A{row},Cartes!$C$2:$C${last},"Héros")',
        5: f"=AVERAGEIF({R},A{row},Cartes!$F$2:$F${last})",
        6: f"=AVERAGEIF({R},A{row},Cartes!$G$2:$G${last})",
        7: f"=AVERAGEIF({R},A{row},Cartes!$H$2:$H${last})",
        8: f"=_xlfn.MINIFS(Cartes!$I$2:$I${last},{R},A{row})",
        9: f"=AVERAGEIF({R},A{row},Cartes!$I$2:$I${last})",
        10: f"=_xlfn.MAXIFS(Cartes!$I$2:$I${last},{R},A{row})",
        11: "–" if i == 0 else f"=H{row}-J{row-1}",
    }
    for j, v in f.items():
        c = sy.cell(row=row, column=j, value=v); c.font = BODY
        c.number_format = "0.0" if j >= 5 else "0"
sy["A10"] = "Total"; sy["A10"].font = B
for j, col in ((2, "B"), (3, "C"), (4, "D")):
    sy.cell(row=10, column=j, value=f"=SUM({col}4:{col}8)").font = B
sy["A12"] = "Marge sur rareté inférieure = puissance min du palier − puissance max du palier précédent. Elle doit rester positive : aucune carte n'est plus forte qu'une carte de rareté supérieure."
sy["A12"].font = Font(name=FONT, italic=True, color="6B6880", size=9)
sy.column_dimensions["A"].width = 14
for c in "BCDEFGHIJK":
    sy.column_dimensions[c].width = 13
sy.row_dimensions[3].height = 32

wb.move_sheet("Synthèse", offset=-1)
wb.save("/home/claude/collectear/Collectear - 300 cartes.xlsx")
print("saved")
for r in RARITIES:
    print(r, [c[0] for c in cards if c[2] == r][:12])
